# -*- coding: utf-8 -*-
"""
Vercel Python Serverless Function: 会話データ(JSON) -> LINE風チャット縦型pptx(バイナリ)
POST /api/generate  body例:
{
  "partner_name": "たなか",
  "partner_status": "オンライン",
  "anchor": "bottom",          // "bottom" | "top" | "center"
  "stagger_ms": 800,
  "conversation": [
     [ ["L","おつかれさま！"], ["R","どうしたの？"] ],   // スライド1
     [ ["R","..."], ["L","..."] ]                         // スライド2
  ]
}
"""
from http.server import BaseHTTPRequestHandler
import json
import math
from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------- 配色・寸法（pptx版と同一） ----------------
FONT = "Hiragino Sans"
EFFECT_MS = 500

C_OUTER   = RGBColor(0xE8, 0xEA, 0xED)
C_BEZEL   = RGBColor(0x0B, 0x0B, 0x0D)
C_CHATBG  = RGBColor(0x8A, 0xAB, 0xD2)
C_HEADER  = RGBColor(0x06, 0xC7, 0x55)
C_SENT    = RGBColor(0x06, 0xC7, 0x55)
C_SENT_TX = RGBColor(0xFF, 0xFF, 0xFF)
C_RECV    = RGBColor(0xFF, 0xFF, 0xFF)
C_RECV_TX = RGBColor(0x1A, 0x1A, 0x1A)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_AVATAR  = RGBColor(0xFF, 0xFF, 0xFF)
C_AVATAR_TX = RGBColor(0x06, 0xC7, 0x55)
C_INPUTBG = RGBColor(0xF4, 0xF5, 0xF7)
C_INPUTFLD= RGBColor(0xFF, 0xFF, 0xFF)
C_MUTE    = RGBColor(0x9A, 0xA0, 0xA6)

SLIDE_W = Inches(7.5)
SLIDE_H = Inches(13.333)
PHONE = dict(x=0.50, y=0.35, w=6.50, h=12.63)
SCREEN = dict(x=PHONE["x"]+0.16, y=PHONE["y"]+0.16,
              w=PHONE["w"]-0.32, h=PHONE["h"]-0.32)
HEADER_H = 1.50
INPUT_H  = 0.85
ISLAND_W, ISLAND_H = 1.70, 0.40

CHAT_PAD = 0.20
CHAT_X = SCREEN["x"] + CHAT_PAD
CHAT_TOP = SCREEN["y"] + HEADER_H + 0.18
CHAT_W = SCREEN["w"] - CHAT_PAD*2
CHAT_BOTTOM = SCREEN["y"] + SCREEN["h"] - INPUT_H - 0.15

BUBBLE_FONT_PT = 15
MAX_BUBBLE_W = 4.05
MIN_BUBBLE_W = 0.70
BUBBLE_PAD_X = 0.14
BUBBLE_PAD_Y = 0.09
LINE_H = BUBBLE_FONT_PT * 1.42 / 72.0
AVATAR_D = 0.56
GAP = 0.20


def set_run_font(run, name=FONT, size=None, color=None, bold=False):
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def add_rect(slide, shape_type, x, y, w, h, fill=None, line=None, line_w=None,
             adj=None, rotation=None, shadow=False):
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1)
    if adj is not None:
        try:
            if isinstance(adj, (list, tuple)):
                for i, v in enumerate(adj):
                    sp.adjustments[i] = v
            else:
                sp.adjustments[0] = adj
        except Exception:
            pass
    if rotation is not None:
        sp.rotation = rotation
    if not shadow:
        sp.shadow.inherit = False
    return sp


def add_text(slide, x, y, w, h, text, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = ln
        set_run_font(r, font, size, color, bold)
    return tb


def disp_units(s):
    u = 0.0
    for ch in s:
        o = ord(ch)
        if o < 0x3000 and o != 0x2026:
            u += 0.5
        else:
            u += 1.0
    return u


def measure_bubble(text):
    char_w = BUBBLE_FONT_PT / 72.0
    cap_w = char_w * 1.06
    cap_per_line = max(1, int((MAX_BUBBLE_W - BUBBLE_PAD_X * 2) / cap_w))
    total_lines = 0
    longest = 0.0
    for seg in text.split("\n"):
        u = disp_units(seg)
        longest = max(longest, u)
        total_lines += max(1, math.ceil(u / cap_per_line))
    if longest <= cap_per_line and "\n" not in text:
        width = longest * char_w + BUBBLE_PAD_X * 2 + 0.18
        width = max(MIN_BUBBLE_W, min(width, MAX_BUBBLE_W))
        lines = 1
    else:
        width = MAX_BUBBLE_W
        lines = total_lines
    height = lines * LINE_H + BUBBLE_PAD_Y * 2 + 0.06
    return width, height, lines


def _soft_shadow(shape):
    spPr = shape._element.spPr
    e = spPr.find(qn("a:effectLst"))
    if e is not None:
        spPr.remove(e)
    effectLst = spPr.makeelement(qn("a:effectLst"), {})
    outer = effectLst.makeelement(qn("a:outerShdw"), {
        "blurRad": "40000", "dist": "20000", "dir": "5400000", "rotWithShape": "0"})
    clr = outer.makeelement(qn("a:srgbClr"), {"val": "000000"})
    alpha = clr.makeelement(qn("a:alpha"), {"val": "18000"})
    clr.append(alpha); outer.append(clr); effectLst.append(outer); spPr.append(effectLst)


def _bubble_tf(bub, text, color):
    tf = bub.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(BUBBLE_PAD_X); tf.margin_right = Inches(BUBBLE_PAD_X)
    tf.margin_top = Inches(BUBBLE_PAD_Y); tf.margin_bottom = Inches(BUBBLE_PAD_Y)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = ln
        set_run_font(r, FONT, BUBBLE_FONT_PT, color)


def add_bubble(slide, side, text, top, partner_initial):
    w, h, _ = measure_bubble(text)
    if side == "L":
        ax = CHAT_X
        av = add_rect(slide, MSO_SHAPE.OVAL, ax, top, AVATAR_D, AVATAR_D, fill=C_AVATAR)
        add_text(slide, ax, top, AVATAR_D, AVATAR_D, partner_initial,
                 13, C_AVATAR_TX, bold=True, align=PP_ALIGN.CENTER)
        bx = ax + AVATAR_D + 0.12
        bub = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, top, w, h, fill=C_RECV)
        bub.adjustments[0] = min(0.5, 0.18 / min(w, h))
        _bubble_tf(bub, text, C_RECV_TX); _soft_shadow(bub)
        spids = [bub.shape_id, av.shape_id]
    else:
        bx = CHAT_X + CHAT_W - w
        bub = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, top, w, h, fill=C_SENT)
        bub.adjustments[0] = min(0.5, 0.18 / min(w, h))
        _bubble_tf(bub, text, C_SENT_TX); _soft_shadow(bub)
        spids = [bub.shape_id]
    return spids, top + h + GAP


def build_chrome(slide, partner_name, partner_status):
    initial = partner_name[0] if partner_name else "?"
    add_rect(slide, MSO_SHAPE.RECTANGLE, 0, 0, 7.5, 13.333, fill=C_OUTER)
    bz = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                  PHONE["x"], PHONE["y"], PHONE["w"], PHONE["h"], fill=C_BEZEL)
    bz.adjustments[0] = 0.5 / min(PHONE["w"], PHONE["h"])
    scr = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                   SCREEN["x"], SCREEN["y"], SCREEN["w"], SCREEN["h"], fill=C_CHATBG)
    scr.adjustments[0] = 0.42 / min(SCREEN["w"], SCREEN["h"])
    hd = add_rect(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
                  SCREEN["x"], SCREEN["y"], SCREEN["w"], HEADER_H, fill=C_HEADER)
    hd.adjustments[0] = 0.42 / min(SCREEN["w"], HEADER_H); hd.adjustments[1] = 0.0
    add_text(slide, SCREEN["x"]+0.45, SCREEN["y"]+0.06, 1.2, 0.40, "9:41",
             14, C_WHITE, bold=True, align=PP_ALIGN.LEFT)
    sbx = SCREEN["x"] + SCREEN["w"] - 1.35
    sby = SCREEN["y"] + 0.16
    for i in range(4):
        bh = 0.07 + i*0.045
        add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                 sbx + i*0.105, sby + (0.20-bh), 0.075, bh, fill=C_WHITE, adj=0.3)
    bxx = SCREEN["x"] + SCREEN["w"] - 0.78
    add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bxx, sby, 0.34, 0.18,
             fill=None, line=C_WHITE, line_w=1.2, adj=0.25)
    add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bxx+0.02, sby+0.025, 0.27, 0.13,
             fill=C_WHITE, adj=0.2)
    add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bxx+0.345, sby+0.05, 0.03, 0.08,
             fill=C_WHITE, adj=0.4)
    isl = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                   SCREEN["x"] + (SCREEN["w"]-ISLAND_W)/2, SCREEN["y"]+0.10,
                   ISLAND_W, ISLAND_H, fill=C_BEZEL)
    isl.adjustments[0] = 0.5
    add_text(slide, SCREEN["x"]+0.30, SCREEN["y"]+0.62, 0.4, 0.7, "‹",
             30, C_WHITE, bold=True, align=PP_ALIGN.LEFT)
    nav_av_x = SCREEN["x"]+0.78; nav_av_y = SCREEN["y"]+0.66
    add_rect(slide, MSO_SHAPE.OVAL, nav_av_x, nav_av_y, 0.62, 0.62, fill=C_WHITE)
    add_text(slide, nav_av_x, nav_av_y, 0.62, 0.62, initial,
             18, C_HEADER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, nav_av_x+0.74, SCREEN["y"]+0.66, 3.0, 0.40, partner_name,
             18, C_WHITE, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.BOTTOM)
    add_text(slide, nav_av_x+0.74, SCREEN["y"]+1.04, 3.0, 0.28, partner_status,
             11, RGBColor(0xDF, 0xF5, 0xE6), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    mx = SCREEN["x"] + SCREEN["w"] - 0.55; my = SCREEN["y"] + 0.92
    for i in range(3):
        add_rect(slide, MSO_SHAPE.OVAL, mx, my + i*0.13, 0.07, 0.07, fill=C_WHITE)
    iby = SCREEN["y"] + SCREEN["h"] - INPUT_H
    ib = add_rect(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
                  SCREEN["x"], iby, SCREEN["w"], INPUT_H, fill=C_INPUTBG, rotation=180)
    ib.adjustments[0] = 0.42 / min(SCREEN["w"], INPUT_H); ib.adjustments[1] = 0.0
    add_text(slide, SCREEN["x"]+0.22, iby+0.18, 0.5, 0.5, "＋",
             20, C_MUTE, align=PP_ALIGN.CENTER)
    fld = add_rect(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                   SCREEN["x"]+0.78, iby+0.18, SCREEN["w"]-1.95, INPUT_H-0.36, fill=C_INPUTFLD)
    fld.adjustments[0] = 0.5
    add_text(slide, SCREEN["x"]+0.95, iby+0.18, 3.0, INPUT_H-0.36,
             "メッセージを入力", 12, C_MUTE, align=PP_ALIGN.LEFT)
    add_rect(slide, MSO_SHAPE.OVAL,
             SCREEN["x"]+SCREEN["w"]-0.92, iby+0.18, INPUT_H-0.36, INPUT_H-0.36, fill=C_HEADER)
    add_text(slide, SCREEN["x"]+SCREEN["w"]-0.92, iby+0.14, INPUT_H-0.36, INPUT_H-0.36,
             "➤", 14, C_WHITE, align=PP_ALIGN.CENTER)


NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" ' \
     'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'


def _effect_par(eid, spid, delay, node_type, grp):
    return f'''<p:par {NS}>
  <p:cTn id="{eid}" presetID="42" presetClass="entr" presetSubtype="8" fill="hold" grpId="{grp}" nodeType="{node_type}">
    <p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>
    <p:childTnLst>
      <p:set>
        <p:cBhvr>
          <p:cTn id="{eid+1}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
          <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:to><p:strVal val="visible"/></p:to>
      </p:set>
      <p:anim calcmode="lin" valueType="num">
        <p:cBhvr additive="base">
          <p:cTn id="{eid+2}" dur="{EFFECT_MS}" fill="hold"/>
          <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
          <p:attrNameLst><p:attrName>ppt_y</p:attrName></p:attrNameLst>
        </p:cBhvr>
        <p:tavLst>
          <p:tav tm="0"><p:val><p:strVal val="ppt_y+0.04"/></p:val></p:tav>
          <p:tav tm="100000"><p:val><p:strVal val="ppt_y"/></p:val></p:tav>
        </p:tavLst>
      </p:anim>
      <p:animEffect transition="in" filter="fade">
        <p:cBhvr>
          <p:cTn id="{eid+3}" dur="{EFFECT_MS}"/>
          <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
        </p:cBhvr>
      </p:animEffect>
    </p:childTnLst>
  </p:cTn>
</p:par>'''


def add_timing(slide, steps, stagger_ms):
    if not steps:
        return
    effects = []; eid = 20; grp = 0; first = True
    for i, spids in enumerate(steps):
        delay = 0 if i == 0 else i * stagger_ms
        for spid in spids:
            node = "clickEffect" if first else "withEffect"
            effects.append(_effect_par(eid, spid, delay, node, grp))
            eid += 4; first = False
        grp += 1
    effects_xml = "\n".join(effects)
    timing_xml = f'''<p:timing {NS}>
  <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
    <p:seq concurrent="1" nextAc="seek">
      <p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
        <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
          <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
            {effects_xml}
          </p:childTnLst></p:cTn></p:par>
        </p:childTnLst></p:cTn></p:par>
      </p:childTnLst></p:cTn>
      <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
      <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
    </p:seq>
  </p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>'''
    slide._element.append(etree.fromstring(timing_xml.encode("utf-8")))


def build_pptx(data):
    """会話データ(dict) -> pptxバイト列"""
    partner_name = (data.get("partner_name") or "相手").strip()
    partner_status = data.get("partner_status", "オンライン")
    anchor = data.get("anchor", "bottom")
    stagger_ms = int(data.get("stagger_ms", 800))
    conversation = data.get("conversation") or []
    initial = partner_name[0] if partner_name else "?"

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for msgs in conversation:
        msgs = [(m[0], m[1]) for m in msgs if len(m) >= 2 and str(m[1]).strip() != ""]
        if not msgs:
            continue
        slide = prs.slides.add_slide(blank)
        build_chrome(slide, partner_name, partner_status)
        heights = [measure_bubble(t)[1] for _, t in msgs]
        block_h = sum(heights) + GAP * (len(msgs) - 1)
        avail = CHAT_BOTTOM - CHAT_TOP
        if block_h > avail:
            start = CHAT_TOP
        elif anchor == "bottom":
            start = CHAT_BOTTOM - block_h
        elif anchor == "center":
            start = CHAT_TOP + (avail - block_h) / 2
        else:
            start = CHAT_TOP
        top = start; steps = []
        for side, text in msgs:
            spids, top = add_bubble(slide, side, text, top, initial)
            steps.append(spids)
        add_timing(slide, steps, stagger_ms)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------- Vercel serverless handler ----------------
class handler(BaseHTTPRequestHandler):
    def _cors(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def do_OPTIONS(self):
        self.send_response(204); self._cors(); self.end_headers()

    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "application/json"); self._cors(); self.end_headers()
        self.wfile.write(b'{"ok":true,"endpoint":"POST /api/generate"}')

    def do_POST(self):
        try:
            length = int(self.headers.get("content-length", 0))
            data = json.loads(self.rfile.read(length) or b"{}")
            if not data.get("conversation"):
                raise ValueError("conversation is empty")
            pptx_bytes = build_pptx(data)
        except Exception as e:
            self.send_response(400)
            self.send_header("Content-Type", "application/json"); self._cors(); self.end_headers()
            self.wfile.write(json.dumps({"error": str(e)}).encode("utf-8"))
            return
        self.send_response(200)
        self.send_header("Content-Type",
                         "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.send_header("Content-Disposition", 'attachment; filename="line_chat.pptx"')
        self.send_header("Content-Length", str(len(pptx_bytes)))
        self._cors(); self.end_headers()
        self.wfile.write(pptx_bytes)
